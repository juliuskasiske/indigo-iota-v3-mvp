"""Read-only database browser for the Control Tower.

Lets the platform owner look at every table — the shared control plane and each
tenant's brain DB — without a psql shell. Strictly read-only: it lists tables,
their columns, and a paged sample of rows.

Two safety rules matter here:
  * **No SQL injection.** Table names never come straight from the request into a
    query string. We first confirm the name exists in ``information_schema`` for
    the target database, then build the statement with ``psycopg.sql.Identifier``.
  * **No secret leakage.** Columns that hold secrets (e.g. ``client_secret``) are
    masked before the rows ever leave this module.
"""
from __future__ import annotations

import csv
import datetime
import decimal
import io
import json
import uuid

from psycopg import sql

from src.db.connection import (
    control_db_name,
    get_control_connection,
    get_tenant_connection,
)

MAX_LIMIT = 200
# Hard cap on a single download so a runaway table can't exhaust memory. Plenty
# for the analytics this surface is for; a bigger pull belongs in a real export.
MAX_EXPORT_ROWS = 100_000

# Column names whose values must never be returned. Matched case-insensitively,
# either exactly or by suffix, so e.g. ``client_secret`` and ``api_secret`` both
# mask but ``field_key`` (ontology) does not.
_SECRET_EXACT = {"password", "secret", "client_secret"}
_SECRET_SUFFIXES = ("_secret", "_password", "_token")


def _is_secret_column(name: str) -> bool:
    low = name.lower()
    return low in _SECRET_EXACT or low.endswith(_SECRET_SUFFIXES)


def _jsonable(value: object) -> object:
    """Coerce a psycopg cell value into something JSON-serializable."""
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, decimal.Decimal):
        return str(value)
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, uuid.UUID):
        return str(value)
    if isinstance(value, (bytes, bytearray, memoryview)):
        return f"<{len(bytes(value))} bytes>"
    if isinstance(value, (dict, list)):
        return value  # jsonb — already JSON-shaped
    return str(value)


def _tenant_db_name(slug: str) -> str:
    with get_control_connection() as conn, conn.cursor() as cur:
        cur.execute(
            """
            SELECT td.db_name
            FROM organizations o
            JOIN tenant_databases td ON td.org_id = o.id
            WHERE o.slug = %s;
            """,
            (slug,),
        )
        row = cur.fetchone()
    if not row:
        raise ValueError(f"No tenant database for slug {slug!r}.")
    return row[0]


def _connect(database_key: str):
    """Open a connection to the database identified by ``database_key``.

    ``control`` -> the control plane; ``tenant:<slug>`` -> that tenant's brain DB.
    """
    if database_key == "control":
        return get_control_connection()
    if database_key.startswith("tenant:"):
        slug = database_key.split(":", 1)[1]
        return get_tenant_connection(_tenant_db_name(slug))
    raise ValueError(f"Unknown database {database_key!r}.")


def list_databases() -> list[dict]:
    """Every database the owner can browse: the control plane + each tenant."""
    out: list[dict] = [
        {
            "key": "control",
            "label": "Control plane",
            "db_name": control_db_name(),
            "kind": "control",
        }
    ]
    with get_control_connection() as conn, conn.cursor() as cur:
        cur.execute(
            """
            SELECT o.slug, o.name, td.db_name
            FROM organizations o
            JOIN tenant_databases td ON td.org_id = o.id
            ORDER BY o.created_at;
            """
        )
        for slug, name, db_name in cur.fetchall():
            out.append(
                {
                    "key": f"tenant:{slug}",
                    "label": f"{name} ({slug})",
                    "db_name": db_name,
                    "kind": "tenant",
                    "slug": slug,
                }
            )
    return out


def list_tables(database_key: str) -> list[dict]:
    """Base tables in the ``public`` schema, each with an exact row count.

    Counts are exact (the data is small); fine for a browse surface.
    """
    with _connect(database_key) as conn, conn.cursor() as cur:
        cur.execute(
            """
            SELECT table_name
            FROM information_schema.tables
            WHERE table_schema = 'public' AND table_type = 'BASE TABLE'
            ORDER BY table_name;
            """
        )
        names = [r[0] for r in cur.fetchall()]
        out: list[dict] = []
        for name in names:
            cur.execute(
                sql.SQL("SELECT COUNT(*) FROM {}").format(sql.Identifier(name))
            )
            out.append({"name": name, "row_count": cur.fetchone()[0]})
    return out


def _order_clause(
    sort_column: str | None, sort_dir: str | None, col_names: list[str]
) -> sql.Composable:
    """ORDER BY for a validated column (else the stable default of column 1).

    ``sort_column`` is only honoured if it's an actual column of this table, so
    a request can never inject an arbitrary expression. Direction is reduced to
    a literal ASC/DESC. NULLS LAST keeps empty values out of the way when sorting.
    """
    if sort_column and sort_column in col_names:
        direction = (
            sql.SQL("DESC") if str(sort_dir).lower() == "desc" else sql.SQL("ASC")
        )
        return sql.SQL("ORDER BY {} {} NULLS LAST").format(
            sql.Identifier(sort_column), direction
        )
    return sql.SQL("ORDER BY 1")


def _where_clause(
    filter_q: str | None, col_names: list[str], secret_idx: set[int]
) -> tuple[sql.Composable, list[object]]:
    """A global, case-insensitive substring filter across every non-secret
    column (each cast to text). Secret columns are excluded so the filter never
    probes masked values. Returns the clause + its bound params."""
    q = (filter_q or "").strip()
    if not q:
        return sql.SQL(""), []
    searchable = [n for i, n in enumerate(col_names) if i not in secret_idx]
    if not searchable:
        return sql.SQL(""), []
    conds = [
        sql.SQL("CAST({} AS text) ILIKE %s").format(sql.Identifier(n))
        for n in searchable
    ]
    clause = sql.SQL("WHERE (") + sql.SQL(" OR ").join(conds) + sql.SQL(")")
    return clause, [f"%{q}%"] * len(searchable)


def _read(
    database_key: str,
    table: str,
    *,
    limit: int,
    offset: int,
    sort_column: str | None,
    sort_dir: str | None,
    filter_q: str | None,
) -> dict:
    """Columns + a window of rows for one table, with optional server-side sort
    and filter. Secrets are masked. Shared by the paged browser and the export."""
    with _connect(database_key) as conn, conn.cursor() as cur:
        # Validate the table exists in this DB BEFORE using its name in a query.
        cur.execute(
            """
            SELECT 1 FROM information_schema.tables
            WHERE table_schema = 'public' AND table_name = %s;
            """,
            (table,),
        )
        if not cur.fetchone():
            raise ValueError(f"No table named {table!r} in this database.")

        cur.execute(
            """
            SELECT column_name, data_type
            FROM information_schema.columns
            WHERE table_schema = 'public' AND table_name = %s
            ORDER BY ordinal_position;
            """,
            (table,),
        )
        col_meta = cur.fetchall()
        columns = [
            {"name": c[0], "type": c[1], "masked": _is_secret_column(c[0])}
            for c in col_meta
        ]
        col_names = [c[0] for c in col_meta]
        secret_idx = {i for i, c in enumerate(col_meta) if _is_secret_column(c[0])}

        where, where_params = _where_clause(filter_q, col_names, secret_idx)

        cur.execute(
            sql.SQL("SELECT COUNT(*) FROM {} ").format(sql.Identifier(table)) + where,
            where_params,
        )
        total = cur.fetchone()[0]

        order = _order_clause(sort_column, sort_dir, col_names)
        cur.execute(
            sql.SQL("SELECT * FROM {} ").format(sql.Identifier(table))
            + where
            + sql.SQL(" ")
            + order
            + sql.SQL(" LIMIT %s OFFSET %s"),
            where_params + [limit, offset],
        )
        raw = cur.fetchall()

    rows: list[list[object]] = []
    for r in raw:
        row: list[object] = []
        for i, val in enumerate(r):
            if i in secret_idx:
                row.append("••••••" if val is not None else None)
            else:
                row.append(_jsonable(val))
        rows.append(row)

    return {
        "table": table,
        "columns": columns,
        "rows": rows,
        "total": total,
        "limit": limit,
        "offset": offset,
    }


def get_rows(
    database_key: str,
    table: str,
    *,
    limit: int = 50,
    offset: int = 0,
    sort_column: str | None = None,
    sort_dir: str | None = "asc",
    filter_q: str | None = None,
) -> dict:
    """Columns + a paged window of rows for one table. Secrets are masked."""
    limit = max(1, min(MAX_LIMIT, limit))
    offset = max(0, offset)
    return _read(
        database_key,
        table,
        limit=limit,
        offset=offset,
        sort_column=sort_column,
        sort_dir=sort_dir,
        filter_q=filter_q,
    )


# ---------------------------------------------------------------------------
# Export (CSV / XLSX) — the full (optionally filtered/sorted) table for offline
# analytics. Same masking as the browser; capped at MAX_EXPORT_ROWS.
# ---------------------------------------------------------------------------

def _cell_str(cell: object) -> str:
    if cell is None:
        return ""
    if isinstance(cell, (dict, list)):
        return json.dumps(cell, ensure_ascii=False)
    return str(cell)


def _to_csv(col_names: list[str], rows: list[list[object]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(col_names)
    for r in rows:
        writer.writerow([_cell_str(c) for c in r])
    # Lead with a BOM so Excel opens the UTF-8 file with the right encoding.
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _sheet_name(table: str) -> str:
    bad = set("[]:*?/\\")
    name = "".join(ch for ch in table if ch not in bad) or "Sheet1"
    return name[:31]


def _to_xlsx(table: str, col_names: list[str], rows: list[list[object]]) -> bytes:
    from openpyxl import Workbook  # lazy: only the export path needs it

    wb = Workbook(write_only=True)  # streams rows, low memory for big tables
    ws = wb.create_sheet(title=_sheet_name(table))
    ws.append(col_names)
    for r in rows:
        ws.append(
            [_cell_str(c) if isinstance(c, (dict, list)) else c for c in r]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _safe_filename(table: str) -> str:
    return (
        "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in table) or "table"
    )


def export_table(
    database_key: str,
    table: str,
    fmt: str,
    *,
    sort_column: str | None = None,
    sort_dir: str | None = "asc",
    filter_q: str | None = None,
) -> tuple[bytes, str, str]:
    """Return (content, media_type, filename) for a CSV/XLSX download of the
    (optionally filtered + sorted) table. Secrets masked; capped at
    MAX_EXPORT_ROWS rows."""
    fmt = (fmt or "csv").lower()
    if fmt not in ("csv", "xlsx"):
        raise ValueError("format must be 'csv' or 'xlsx'.")

    data = _read(
        database_key,
        table,
        limit=MAX_EXPORT_ROWS,
        offset=0,
        sort_column=sort_column,
        sort_dir=sort_dir,
        filter_q=filter_q,
    )
    col_names = [c["name"] for c in data["columns"]]
    rows = data["rows"]
    stem = _safe_filename(table)

    if fmt == "csv":
        return _to_csv(col_names, rows), "text/csv; charset=utf-8", f"{stem}.csv"
    return (
        _to_xlsx(table, col_names, rows),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        f"{stem}.xlsx",
    )
